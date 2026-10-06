"""Genera l'Excel di riepilogo: reports/riepilogo_<stagione>.xlsx

    python -m basket.export_excel
"""

import sys

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import analysis as A
from . import config, db
from . import views as V

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TWO_DECIMALS = ("rate", "FTA/FGA", "Variabilità")


def build_sheets(conn) -> dict[str, pd.DataFrame]:
    bs = A.load_box_squadra(conn)
    bg = A.load_box_giocatore(conn)
    teams = A.team_summary(bs).sort_values(["campionato_id", "squadra"])
    players = A.player_summary(bg, bs).sort_values(["campionato_id", "punti_pg"],
                                                   ascending=[True, False])
    profilo = A.team_profile(bs, bg, A.load_quarters(conn)).sort_values(
        ["campionato_id", "squadra"])
    trend = A.player_trend(bg, bs).sort_values(["campionato_id", "valutazione_pg_delta"],
                                               ascending=[True, False])
    return {
        "Classifica": V.select(A.standings(bs), V.CLASSIFICA),
        "Squadre medie": V.select(teams, V.TEAM_MEDIE),
        "Squadre per 40": V.select(teams, V.TEAM_P40),
        "Squadre avanzate": V.select(teams, V.TEAM_AVANZATE),
        "Squadre profilo": V.select(profilo, V.TEAM_PROFILO),
        "Squadre casa-trasf.": V.casa_trasferta_squadre(bs),
        "Giocatori medie": V.select(players, V.PLAYER_MEDIE),
        "Giocatori totali": V.select(players, V.PLAYER_TOTALI),
        "Giocatori per 40": V.select(players, V.PLAYER_P40),
        "Giocatori avanzate": V.select(players, V.PLAYER_AVANZATE),
        "Giocatori ruolo": V.select(players, V.PLAYER_RUOLO),
        "Giocatori trend 5": V.select(trend, V.TREND),
        "Giocatori casa-trasf.": V.casa_trasferta_giocatori(bg, bs),
        "Partite": V.partite(conn),
    }


def write_excel(sheets: dict[str, pd.DataFrame], path):
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "C2" if name.startswith("Giocatori") else "B2"
            if ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.fill, cell.font = HEADER_FILL, HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for i, col in enumerate(df.columns, start=1):
                letter = get_column_letter(i)
                is_float = pd.api.types.is_float_dtype(df[col])
                width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)]) + 2
                ws.column_dimensions[letter].width = min(max(width if not is_float else 8, 6), 38)
                if is_float:
                    fmt = "0.00" if any(k in str(col) for k in TWO_DECIMALS) else "0.0"
                    for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                        cell.number_format = fmt


def main() -> int:
    conn = db.connect(config.DB_PATH, readonly=True)
    sheets = build_sheets(conn)
    config.REPORTS_DIR.mkdir(exist_ok=True)
    out = config.REPORTS_DIR / f"riepilogo_{config.STAGIONE}.xlsx"
    write_excel(sheets, out)
    print(f"Scritto {out} ({', '.join(f'{k}: {len(v)}' for k, v in sheets.items())})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
