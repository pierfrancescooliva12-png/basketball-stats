"""Report prepartita in Word (.docx), modificabile dallo staff: visione di squadra, scouting dei
giocatori con spazio per le note, statistiche dei giocatori. Stessi contenuti delle prime
pagine del PDF (basket.report_pdf)."""

import io
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import config, scouting
from . import scheda_giocatori as SG
from . import sintesi as SI

ARANCIO = RGBColor(0xFF, 0x7A, 0x1A)
BLU = RGBColor(0x2A, 0x78, 0xD6)
ROSSO = RGBColor(0xD9, 0x59, 0x26)
GRIGIO = RGBColor(0x6B, 0x74, 0x89)
INK = RGBColor(0x1A, 0x22, 0x33)
NAVY = RGBColor(0x0B, 0x12, 0x20)
TINTA = {"forza": BLU, "debole": ROSSO, "media": GRIGIO, "neutro": INK}


def _it(v, d=1) -> str:
    return SG._it(v, d)


def _testo(par, testo, bold=False, size=None, colore=None, italic=False):
    run = par.add_run(testo)
    run.bold, run.italic = bold, italic
    if size:
        run.font.size = Pt(size)
    if colore is not None:
        run.font.color.rgb = colore
    return run


def _sfondo(cella, hex_colore: str):
    tc = cella._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_colore)
    tc.append(shd)


def _titolo(doc, testo, livello=1):
    h = doc.add_heading(level=livello)
    r = _testo(h, testo.upper() if livello == 1 else testo, colore=ARANCIO if livello == 1 else NAVY)
    r.italic = True
    return h


def _fascia(doc, occhiello, titolo, sotto):
    """Fascia scura a tutta larghezza in testa alla pagina."""
    t = doc.add_table(rows=1, cols=1)
    cella = t.rows[0].cells[0]
    _sfondo(cella, "0B1220")
    p = cella.paragraphs[0]
    _testo(p, occhiello, bold=True, size=8, colore=ARANCIO)
    p = cella.add_paragraph()
    r = _testo(p, titolo.upper(), bold=True, size=26, colore=RGBColor(0xFF, 0xFF, 0xFF))
    r.italic = True
    if sotto:
        _testo(cella.add_paragraph(), sotto, size=9, colore=RGBColor(0xC9, 0xD1, 0xE0))
    riga = doc.add_paragraph()
    _testo(riga, " ", size=2)
    bordo = OxmlElement("w:pBdr")
    sotto_b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "24"), ("w:space", "1"), ("w:color", "FF7A1A")):
        sotto_b.set(qn(k), v)
    bordo.append(sotto_b)
    riga._p.get_or_add_pPr().append(bordo)


def _elenco(doc, voci, segno="•", colore=ARANCIO, vuoto="Nessuna indicazione marcata."):
    if not voci:
        p = doc.add_paragraph()
        _testo(p, vuoto, italic=True, colore=GRIGIO, size=9)
        return
    for v in voci:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        _testo(p, f"{segno} ", bold=True, colore=colore)
        _testo(p, v, size=10)


def _spazio_note(doc, righe=3, titolo="Note dello staff"):
    p = doc.add_paragraph()
    _testo(p, titolo, bold=True, size=9, colore=GRIGIO)
    for _ in range(righe):
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(0)
        _testo(q, "_" * 95, colore=RGBColor(0xC9, 0xD1, 0xE0), size=9)


def build_docx(ctx, squadra_id: int, mia: int | None = None, foto: bool = True) -> bytes:
    """Report prepartita su squadra_id in formato Word."""
    r = scouting.build_report(ctx.conn, squadra_id, ctx.stagione)
    s = SI.squadra(ctx, squadra_id, mia, r)
    giocatori = SG.schede(ctx, squadra_id, foto=foto)
    doc = Document()
    sez = doc.sections[0]
    sez.left_margin = sez.right_margin = Cm(1.6)
    sez.top_margin = sez.bottom_margin = Cm(1.5)
    stile = doc.styles["Normal"]
    stile.font.name, stile.font.size = "Calibri", Pt(10)

    # intestazione "Match day": fascia scura con occhiello arancione e nome in corsivo
    ng = ctx.next_game(squadra_id)
    meta = ""
    if ng:
        nomi = ctx.nomi_squadre
        meta = (f"{ng['giornata']}ª giornata · {ng['data'][8:10]}/{ng['data'][5:7]} {ng.get('ora') or ''} · "
                f"{nomi.get(ng['casa_id'], '')} – {nomi.get(ng['ospite_id'], '')}")
    _fascia(doc, f"SCOUTING REPORT · {r.campionato.upper()} {config.STAGIONI.get(ctx.stagione, ctx.stagione)}",
            r.squadra, f"{r.record}  ·  {r.posizione}° posto" + (f"  ·  {meta}" if meta else ""))

    # visione di squadra
    _titolo(doc, "La squadra in 30 secondi")
    p = doc.add_paragraph()
    _testo(p, s["sintesi"], size=11)
    _titolo(doc, "Come stanno in campionato", 2)
    voci = SI.semaforo(r.profilo)
    tab = doc.add_table(rows=1, cols=4)
    tab.style = "Light Grid Accent 1"
    for c, h in zip(tab.rows[0].cells, ("Voce", "Valore", "Posizione", "Giudizio")):
        c.text = h
    for v in voci:
        c = tab.add_row().cells
        c[0].text = f"{v['etichetta']} ({v['spiega']})"
        c[1].text = _it(v["valore"])
        c[2].text = f"{v['posizione']}° su {v['squadre']}"
        _testo(c[3].paragraphs[0], v["giudizio"], bold=True, colore=TINTA[v["tipo"]])
    for titolo, chiave, colore in (("In difesa · contro il loro attacco", "in_difesa", BLU),
                                   ("In attacco · contro la loro difesa", "in_attacco", ARANCIO),
                                   ("Ritmo, momenti e finali", "ritmo", GRIGIO)):
        _titolo(doc, titolo, 2)
        _elenco(doc, s[chiave], "›", colore)
    if s["da_ricordare"]:
        _titolo(doc, "Da ricordare", 2)
        for i, x in enumerate(s["da_ricordare"][:3], start=1):
            p = doc.add_paragraph()
            _testo(p, f"{i}  ", bold=True, size=14, colore=ARANCIO)
            _testo(p, x, bold=True, size=11)
    if s["attenzione"]:
        p = doc.add_paragraph()
        _testo(p, f"⚠ {s['attenzione']}", italic=True, size=8, colore=GRIGIO)
    _spazio_note(doc, 5, "Note dello staff sulla squadra")

    # giocatori
    if giocatori:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _fascia(doc, "SCOUTING DEI GIOCATORI", r.squadra, "")
        p = doc.add_paragraph()
        _testo(p, "Giocatori di rotazione (almeno 10 minuti a partita), in ordine di minuti. "
                  "+ cosa fa meglio, − cosa fa peggio in modo evidente, rispetto ai giocatori di "
                  "rotazione del campionato.", size=8, colore=GRIGIO, italic=True)
        for g in giocatori:
            _giocatore(doc, g)

        # statistiche
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _titolo(doc, "Statistiche dei giocatori", 2)
        pl = ctx.players[ctx.players["squadra_id"] == squadra_id].sort_values("minuti", ascending=False).head(12)
        head = ["Giocatore", "PG", "Min", "Pt", "Rimb", "Ass", "Perse", "3PM-A", "TL%", "TS%"]
        tab = doc.add_table(rows=1, cols=len(head))
        tab.style = "Light Grid Accent 1"
        for c, h in zip(tab.rows[0].cells, head):
            c.text = h
        for x in pl.itertuples():
            vals = [x.giocatore, str(x.partite), _it(x.minuti_pg), _it(x.punti_pg), _it(x.rimb_tot_pg),
                    _it(x.assist_pg), _it(x.perse_pg), f"{_it(x.t3m_pg)}-{_it(x.t3a_pg)}",
                    _it(x.tl_pct, 0), _it(x.ts_pct)]
            for c, v in zip(tab.add_row().cells, vals):
                c.text = v
        for row in tab.rows:
            for c in row.cells:
                for par in c.paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(8.5)

    p = sez.footer.paragraphs[0]
    _testo(p, f"{config.BRAND} · {r.squadra} · "
              f"{date.today().strftime('%d/%m/%Y')}", size=7, colore=GRIGIO)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _giocatore(doc, g: dict):
    tab = doc.add_table(rows=1, cols=2)
    tab.alignment = WD_TABLE_ALIGNMENT.LEFT
    tab.autofit = False
    sx, dx = tab.rows[0].cells
    for col, cella, w in ((tab.columns[0], sx, Cm(3.2)), (tab.columns[1], dx, Cm(14.6))):
        col.width = cella.width = w
    _sfondo(sx, "0B1220")
    _sfondo(dx, "FAFBFD")
    p = sx.paragraphs[0]
    if g.get("maglia"):
        _testo(p, f"#{g['maglia']}  ", bold=True, size=12, colore=ARANCIO)
        p = sx.add_paragraph()
    if g.get("foto"):
        try:
            p.add_run().add_picture(io.BytesIO(g["foto"]), width=Cm(2.8))
        except Exception:
            _testo(p, "".join(x[0] for x in g["nome"].split()[:2]).upper(), bold=True, size=20, colore=RGBColor(0x5F, 0x73, 0x97))
    else:
        _testo(p, "".join(x[0] for x in g["nome"].split()[:2]).upper(), bold=True, size=20, colore=RGBColor(0x5F, 0x73, 0x97))
    p = dx.paragraphs[0]
    r_nome = _testo(p, g["nome"], bold=True, size=13, colore=NAVY)
    r_nome.italic = True
    meta = [x for x in (g.get("ruolo"),
                        f"{g['altezza']} cm" if g.get("altezza") else None,
                        f"{g['eta']} anni" if g.get("eta") else None, g.get("nazionalita")) if x]
    if meta:
        _testo(dx.add_paragraph(), " · ".join(meta), size=8.5, colore=GRIGIO)
    _testo(dx.add_paragraph(), "   ".join(f"{k} {v}" for k, v in g["numeri"]), bold=True, size=10)
    _testo(dx.add_paragraph(), g["tiro"], size=8.5, colore=GRIGIO)
    for segno, voci, colore in (("+", g["forza"], BLU), ("−", g["deboli"], ROSSO)):
        for v in voci:
            q = dx.add_paragraph()
            q.paragraph_format.space_after = Pt(0)
            _testo(q, f"{segno} ", bold=True, colore=colore)
            _testo(q, v, size=9.5)
    _spazio_note(doc, 2)
